#!/usr/bin/env python3
r"""Create the four main-paper tables and six publication figures for HEAPO.

Run from the repository root on Windows:

    py -3.11 .\scripts\make_manuscript_outputs.py \
        --run-dir .\outputs\heapo_analysis

The script does not retrain any model. It reads the completed run outputs and
writes:

    outputs/heapo_analysis/manuscript_outputs/
        tables/HEAPO_main_tables.xlsx
        tables/Table_1_*.csv ... Table_4_*.csv
        figures/Figure_1_*.png/.pdf ... Figure_6_*.png/.pdf
        MANUSCRIPT_CAPTIONS.md
        manuscript_output_manifest.json

Required packages:

    py -3.11 -m pip install numpy pandas pyarrow matplotlib openpyxl

The primary inferential inputs are the household-cluster bootstrap tables.
Nested household/day bootstrap results are used only to flag robustness.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Sequence

import matplotlib

matplotlib.use("Agg")
import matplotlib.colors as mcolors
import matplotlib.pyplot as plt
from matplotlib.lines import Line2D
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch
import numpy as np
import pandas as pd


HORIZONS = [1, 4, 24, 96]
HORIZON_LABELS = {1: "15 min", 4: "1 h", 24: "6 h", 96: "24 h"}
MODEL_LABELS = {
    "persistence": "Persistence",
    "seasonal_daily": "Daily seasonal",
    "seasonal_weekly": "Weekly seasonal",
    "elastic_net": "Elastic Net",
    "lightgbm": "LightGBM",
    "tft": "TFT",
    "perfect_foresight": "Perfect foresight",
}
MODEL_ORDER = [
    "persistence",
    "seasonal_daily",
    "seasonal_weekly",
    "elastic_net",
    "lightgbm",
    "tft",
]
STRATUM_ORDER = [
    ("seen_test", "ASHP"),
    ("seen_test", "GSHP"),
    ("unseen_test", "ASHP"),
    ("unseen_test", "GSHP"),
]
STRATUM_LABELS = {
    ("seen_test", "ASHP"): "Seen households — ASHP",
    ("seen_test", "GSHP"): "Seen households — GSHP",
    ("unseen_test", "ASHP"): "Unseen households — ASHP",
    ("unseen_test", "GSHP"): "Unseen households — GSHP",
}

# Okabe-Ito-inspired, color-blind-safe palette.
COLORS = {
    "persistence": "#7A7A7A",
    "seasonal_daily": "#009E73",
    "seasonal_weekly": "#56B4E9",
    "elastic_net": "#CC79A7",
    "lightgbm": "#0072B2",
    "tft": "#D55E00",
    "perfect_foresight": "#222222",
}
BLUE = "#1F4E79"
GREEN = "#2E7D32"
AMBER = "#E69F00"
RED = "#C43C39"
GRID = "#D9DEE5"


@dataclass(frozen=True)
class Sources:
    audit: Path
    split_manifest: Path
    common_metrics: Path
    common_intervals: Path
    tabular_primary: Path
    tabular_nested: Path
    tft_primary: Path
    tft_nested: Path
    shap: Path
    portfolio_summary: Path
    portfolio_detail: Path
    portfolio_corrections: Path


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Generate the HEAPO main-paper tables and figures without retraining."
    )
    parser.add_argument(
        "--run-dir",
        type=Path,
        default=Path("outputs/heapo_analysis"),
        help="Completed canonical run directory.",
    )
    parser.add_argument(
        "--out-dir",
        type=Path,
        default=None,
        help="Output directory (default: RUN_DIR/manuscript_outputs).",
    )
    parser.add_argument(
        "--statistics-dir",
        type=Path,
        default=None,
        help="Optional directory containing final household-cluster statistics; default RUN_DIR.",
    )
    parser.add_argument(
        "--audit-dir",
        type=Path,
        default=None,
        help="Optional audit directory; default RUN_DIR/audit.",
    )
    parser.add_argument(
        "--shap-file",
        type=Path,
        default=None,
        help="Optional explicit SHAP importance CSV/Parquet file.",
    )
    parser.add_argument("--dpi", type=int, default=600, help="PNG resolution (default: 600 dpi).")
    parser.add_argument(
        "--tiff",
        action="store_true",
        help="Also save LZW-compressed TIFF files in addition to PNG and PDF.",
    )
    return parser.parse_args()


def configure_style() -> None:
    plt.rcParams.update(
        {
            "font.family": "DejaVu Sans",
            "font.size": 9.0,
            "axes.titlesize": 10.5,
            "axes.labelsize": 9.5,
            "axes.titleweight": "bold",
            "axes.edgecolor": "#4C5664",
            "axes.linewidth": 0.8,
            "xtick.labelsize": 8.5,
            "ytick.labelsize": 8.5,
            "legend.fontsize": 8.2,
            "figure.dpi": 120,
            "savefig.facecolor": "white",
            "pdf.fonttype": 42,
            "ps.fonttype": 42,
        }
    )


def first_existing(candidates: Sequence[Path], label: str) -> Path:
    for path in candidates:
        if path and path.exists():
            return path.resolve()
    attempted = "\n  - ".join(str(x) for x in candidates if x)
    raise FileNotFoundError(f"Missing {label}. Tried:\n  - {attempted}")


def find_shap_file(run_dir: Path, explicit: Path | None) -> Path:
    if explicit is not None:
        return first_existing([explicit], "SHAP importance file")
    preferred = [
        run_dir / "shap" / "shap_importance.parquet",
        run_dir / "shap" / "shap_importance.csv",
        run_dir / "models" / "shap_importance.parquet",
        run_dir / "models" / "shap_importance.csv",
        run_dir / "metrics" / "shap_importance.parquet",
        run_dir / "metrics" / "shap_importance.csv",
    ]
    for path in preferred:
        if path.exists():
            return path.resolve()
    matches = sorted(
        p
        for p in run_dir.rglob("*")
        if p.is_file()
        and "shap" in p.name.lower()
        and "importance" in p.name.lower()
        and p.suffix.lower() in {".csv", ".parquet"}
        and "manuscript_outputs" not in p.parts
    )
    if matches:
        return matches[0].resolve()
    raise FileNotFoundError(
        "SHAP importance file was not found. Supply it explicitly, for example:\n"
        "  --shap-file .\\outputs\\heapo_analysis\\shap\\shap_importance.csv"
    )


def resolve_sources(
    run_dir: Path,
    statistics_dir: Path | None,
    audit_dir: Path | None,
    shap_file: Path | None,
) -> Sources:
    run_dir = run_dir.resolve()
    stats = (statistics_dir or run_dir).resolve()
    audit_root = (audit_dir or run_dir / "audit").resolve()
    metrics = run_dir / "metrics"
    stats_metrics = stats / "metrics" if (stats / "metrics").exists() else stats
    portfolios = run_dir / "portfolios"
    stats_portfolios = stats / "portfolios" if (stats / "portfolios").exists() else stats

    return Sources(
        audit=first_existing(
            [audit_root / "household_audit.parquet", run_dir / "audit" / "household_audit.parquet"],
            "household audit",
        ),
        split_manifest=first_existing(
            [run_dir / "splits" / "split_manifest.json"], "split manifest"
        ),
        common_metrics=first_existing(
            [metrics / "common_index_model_metrics.parquet"], "common-index model metrics"
        ),
        common_intervals=first_existing(
            [
                stats_metrics / "common_index_household_cluster_metric_intervals.parquet",
                metrics / "common_index_household_cluster_metric_intervals.parquet",
                metrics / "common_index_bootstrap_intervals.parquet",
            ],
            "common-index bootstrap intervals",
        ),
        tabular_primary=first_existing(
            [
                stats_metrics / "household_cluster_paired_tests.parquet",
                metrics / "household_cluster_paired_tests.parquet",
            ],
            "household-cluster tabular paired tests",
        ),
        tabular_nested=first_existing(
            [
                stats_metrics / "nested_household_day_paired_tests.parquet",
                metrics / "nested_household_day_paired_tests.parquet",
            ],
            "nested tabular paired tests",
        ),
        tft_primary=first_existing(
            [
                stats_metrics / "common_index_household_cluster_tft_paired_tests.parquet",
                metrics / "common_index_household_cluster_tft_paired_tests.parquet",
            ],
            "household-cluster TFT paired tests",
        ),
        tft_nested=first_existing(
            [
                stats_metrics / "common_index_nested_household_day_tft_paired_tests.parquet",
                metrics / "common_index_nested_household_day_tft_paired_tests.parquet",
            ],
            "nested TFT paired tests",
        ),
        shap=find_shap_file(run_dir, shap_file),
        portfolio_summary=first_existing(
            [
                portfolios / "peak_management_summary.parquet",
            ],
            "peak-management summary",
        ),
        portfolio_detail=first_existing(
            [portfolios / "peak_management_detail.parquet"],
            "peak-management origin detail",
        ),
        portfolio_corrections=first_existing(
            [
                portfolios / "portfolio_interval_corrections.parquet",
                stats_portfolios / "portfolio_interval_corrections.parquet",
            ],
            "portfolio interval corrections",
        ),
    )


def read_frame(path: Path) -> pd.DataFrame:
    if path.suffix.lower() == ".csv":
        return pd.read_csv(path)
    try:
        return pd.read_parquet(path)
    except ImportError as exc:
        # Allows validation beside the supplied minimal decoder; normal users
        # should use pyarrow, as listed in the installation command.
        try:
            from parquet_minireader import read_parquet  # type: ignore

            return read_parquet(path)
        except Exception:
            raise ImportError(
                "Parquet support is missing. Install it with: "
                "py -3.11 -m pip install pyarrow"
            ) from exc


def require_columns(frame: pd.DataFrame, columns: Iterable[str], source: Path) -> None:
    missing = sorted(set(columns) - set(frame.columns))
    if missing:
        raise ValueError(f"{source} is missing required columns: {missing}")


def panel_label(ax: plt.Axes, label: str) -> None:
    ax.text(
        0.02,
        0.98,
        label,
        transform=ax.transAxes,
        fontsize=9.2,
        fontweight="bold",
        va="top",
        ha="left",
        color="#111111",
        bbox={"facecolor": "white", "edgecolor": "none", "alpha": 0.82, "pad": 1.5},
        zorder=10,
    )


def panel_header(ax: plt.Axes, label: str) -> None:
    """Place a compact panel identifier immediately above the plotting area."""
    ax.text(
        0.0,
        1.025,
        label,
        transform=ax.transAxes,
        fontsize=9.2,
        fontweight="bold",
        va="bottom",
        ha="left",
        color="#111111",
        clip_on=False,
        zorder=10,
    )


def save_figure(fig: plt.Figure, stem: Path, dpi: int, tiff: bool) -> list[Path]:
    stem.parent.mkdir(parents=True, exist_ok=True)
    outputs = []
    png = stem.with_suffix(".png")
    pdf = stem.with_suffix(".pdf")
    # Write raster outputs atomically and verify the complete image stream.
    # This avoids leaving a truncated PNG if a filesystem write is interrupted.
    from PIL import Image

    temporary_png = png.with_name(f"{png.stem}.tmp.png")
    last_error: Exception | None = None
    for _attempt in range(2):
        try:
            fig.canvas.draw()
            fig.savefig(
                temporary_png,
                format="png",
                dpi=dpi,
                bbox_inches="tight",
                pad_inches=0.04,
            )
            with Image.open(temporary_png) as image:
                image.verify()
            temporary_png.replace(png)
            last_error = None
            break
        except Exception as exc:
            last_error = exc
            if temporary_png.exists():
                temporary_png.unlink()
    if last_error is not None:
        raise RuntimeError(f"Could not write a complete PNG file: {png}") from last_error
    fig.savefig(pdf, bbox_inches="tight", pad_inches=0.04)
    outputs.extend([png, pdf])
    if tiff:
        tif = stem.with_suffix(".tiff")
        fig.savefig(tif, dpi=dpi, bbox_inches="tight", pad_inches=0.04, pil_kwargs={"compression": "tiff_lzw"})
        outputs.append(tif)
    # Defensive cleanup for interrupted/retried raster writes.
    temporary_png.unlink(missing_ok=True)
    plt.close(fig)
    return outputs


def hp_counts(ids: Sequence[int], mapping: dict[int, str]) -> tuple[int, int, int]:
    values = [mapping.get(int(value), "Unknown") for value in ids]
    return len(values), values.count("ASHP"), values.count("GSHP")


def build_table_1(audit: pd.DataFrame, split_manifest: dict, intervals: pd.DataFrame) -> pd.DataFrame:
    require_columns(audit, ["household_id", "hp_type", "eligible"], Path("household_audit"))
    audit = audit.copy()
    audit["household_id"] = pd.to_numeric(audit["household_id"], errors="raise").astype(int)
    audit["hp_type"] = audit["hp_type"].fillna("Unknown")
    mapping = dict(zip(audit.household_id, audit.hp_type))

    all_counts = audit.hp_type.value_counts()
    eligible = audit[audit.eligible.astype(bool)]
    eligible_counts = eligible.hp_type.value_counts()
    train = hp_counts(split_manifest["training_households"], mapping)
    validation = hp_counts(split_manifest["unseen_validation_households"], mapping)
    unseen = hp_counts(split_manifest["unseen_test_households"], mapping)

    seen_counts = {"ASHP": 0, "GSHP": 0}
    if "n_households" in intervals.columns:
        seen = intervals[intervals["split"] == "seen_test"]
        for hp in seen_counts:
            values = seen.loc[seen.HP_Type == hp, "n_households"]
            if not values.empty:
                seen_counts[hp] = int(values.max())

    rows = [
        {
            "Cohort": "Complete HEAPO dataset",
            "Total": int(len(audit)),
            "ASHP": int(all_counts.get("ASHP", 0)),
            "GSHP": int(all_counts.get("GSHP", 0)),
            "Unknown HP type": int(all_counts.get("Unknown", 0)),
            "Role": "Source population",
        },
        {
            "Cohort": "Eligible households",
            "Total": int(len(eligible)),
            "ASHP": int(eligible_counts.get("ASHP", 0)),
            "GSHP": int(eligible_counts.get("GSHP", 0)),
            "Unknown HP type": 0,
            "Role": "Analysis cohort",
        },
        {
            "Cohort": "Training households",
            "Total": train[0],
            "ASHP": train[1],
            "GSHP": train[2],
            "Unknown HP type": 0,
            "Role": "Model development",
        },
        {
            "Cohort": "Unseen validation households",
            "Total": validation[0],
            "ASHP": validation[1],
            "GSHP": validation[2],
            "Unknown HP type": 0,
            "Role": "Household-level validation",
        },
        {
            "Cohort": "Unseen test households",
            "Total": unseen[0],
            "ASHP": unseen[1],
            "GSHP": unseen[2],
            "Unknown HP type": 0,
            "Role": "Final generalization test",
        },
        {
            "Cohort": "Independent households in seen test",
            "Total": int(seen_counts["ASHP"] + seen_counts["GSHP"]),
            "ASHP": int(seen_counts["ASHP"]),
            "GSHP": int(seen_counts["GSHP"]),
            "Unknown HP type": 0,
            "Role": "Primary within-household inference",
        },
    ]
    return pd.DataFrame(rows)


def weighted_model_summary(common_metrics: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    required = [
        "model", "split", "HP_Type", "horizon", "n", "mae", "mase",
        "picp", "mean_interval_width", "wis", "approx_crps",
    ]
    require_columns(common_metrics, required, Path("common_index_model_metrics"))
    data = common_metrics[
        common_metrics.horizon.isin(HORIZONS)
        & common_metrics["split"].isin(["seen_test", "unseen_test"])
    ].copy()
    rows = []
    for (split, hp, model), group in data.groupby(["split", "HP_Type", "model"], observed=True):
        weights = pd.to_numeric(group.n, errors="coerce").fillna(0).to_numpy(float)
        if weights.sum() <= 0:
            continue
        row = {"split": split, "HP_Type": hp, "model": model, "n_weight": int(weights.sum())}
        for metric in [
            "mae", "mase", "picp", "mean_interval_width", "wis", "approx_crps"
        ]:
            values = pd.to_numeric(group[metric], errors="coerce").to_numpy(float)
            valid = np.isfinite(values) & np.isfinite(weights) & (weights > 0)
            row[metric] = float(np.average(values[valid], weights=weights[valid]))
        rows.append(row)
    weighted = pd.DataFrame(rows)

    summary_rows = []
    for split, hp in STRATUM_ORDER:
        group = weighted[(weighted["split"] == split) & (weighted.HP_Type == hp)].copy()
        if group.empty:
            continue
        point = group.loc[group.mae.idxmin()]
        probabilistic = group.loc[group.approx_crps.idxmin()]
        lgb = group[group.model == "lightgbm"].iloc[0]
        tft = group[group.model == "tft"].iloc[0]
        summary_rows.append(
            {
                "Evaluation stratum": STRATUM_LABELS[(split, hp)],
                "Point-forecast winner": MODEL_LABELS.get(point.model, point.model),
                "Winner weighted MAE": point.mae,
                "LightGBM weighted MAE": lgb.mae,
                "TFT weighted MAE": tft.mae,
                "Probabilistic winner": MODEL_LABELS.get(probabilistic.model, probabilistic.model),
                "Winner weighted approx. CRPS": probabilistic.approx_crps,
                "Probabilistic-winner PICP": probabilistic.picp,
                "Probabilistic-winner interval width": probabilistic.mean_interval_width,
                "Probabilistic-winner WIS": probabilistic.wis,
            }
        )
    return pd.DataFrame(summary_rows), weighted


def merge_paired(primary: pd.DataFrame, nested: pd.DataFrame) -> pd.DataFrame:
    keys = ["split", "HP_Type", "horizon", "candidate", "comparator"]
    require_columns(
        primary,
        keys + ["n_households", "mae_delta", "ci_low", "ci_high", "holm_p_value", "confirmatory_holm"],
        Path("primary paired tests"),
    )
    require_columns(
        nested,
        keys + ["mae_delta", "ci_low", "ci_high", "holm_p_value", "confirmatory_holm"],
        Path("nested paired tests"),
    )
    return primary.merge(nested, on=keys, how="left", suffixes=("_primary", "_nested"), validate="one_to_one")


def build_table_3(tabular: pd.DataFrame, tft: pd.DataFrame) -> pd.DataFrame:
    frames = []
    for family, frame in [("Tabular", tabular), ("TFT", tft)]:
        selected = frame[frame.confirmatory_holm_primary.astype(bool)].copy()
        selected["Model family"] = family
        frames.append(selected)
    combined = pd.concat(frames, ignore_index=True)
    combined["family_order"] = combined["Model family"].map({"Tabular": 0, "TFT": 1})
    combined["horizon_order"] = combined.horizon.map({h: i for i, h in enumerate(HORIZONS)})
    combined["comparator_order"] = combined.comparator.map({m: i for i, m in enumerate(MODEL_ORDER)})
    combined = combined.sort_values(["family_order", "horizon_order", "comparator_order"])
    output = pd.DataFrame(
        {
            "Model family": combined["Model family"],
            "Evaluation stratum": [STRATUM_LABELS.get((s, h), f"{s} — {h}") for s, h in zip(combined["split"], combined.HP_Type)],
            "Horizon": combined.horizon.map(HORIZON_LABELS),
            "Candidate": combined.candidate.map(lambda x: MODEL_LABELS.get(x, x)),
            "Comparator": combined.comparator.map(lambda x: MODEL_LABELS.get(x, x)),
            "Independent households": combined.n_households_primary.astype(int),
            "Primary delta MAE": combined.mae_delta_primary,
            "Primary 95% CI low": combined.ci_low_primary,
            "Primary 95% CI high": combined.ci_high_primary,
            "Primary Holm p": combined.holm_p_value_primary,
            "Nested delta MAE": combined.mae_delta_nested,
            "Nested 95% CI low": combined.ci_low_nested,
            "Nested 95% CI high": combined.ci_high_nested,
            "Nested Holm p": combined.holm_p_value_nested,
            "Robust in both designs": combined.confirmatory_holm_nested.fillna(False).astype(bool),
            "Direction": np.where(combined.mae_delta_primary < 0, "Candidate better", "Candidate worse"),
        }
    )
    return output.reset_index(drop=True)


def valid_portfolio_pairs(corrections: pd.DataFrame) -> set[tuple[str, int]]:
    require_columns(
        corrections,
        ["portfolio_type", "portfolio_size", "calibration_valid"],
        Path("portfolio interval corrections"),
    )
    grouped = corrections.groupby(["portfolio_type", "portfolio_size"], observed=True).calibration_valid.all()
    return {(str(hp), int(size)) for (hp, size), valid in grouped.items() if bool(valid)}


def filter_portfolio_scenario(
    portfolio: pd.DataFrame,
    valid_pairs: set[tuple[str, int]],
) -> pd.DataFrame:
    """Select the pre-specified scenario and calibration-valid strata."""
    data = portfolio.copy()
    data["portfolio_size"] = pd.to_numeric(data.portfolio_size, errors="coerce").astype("Int64")
    data = data[
        (np.isclose(pd.to_numeric(data.flexible_fraction), 0.20))
        & (pd.to_numeric(data.shift_hours) == 3)
        & data.strategy.isin(["median", "conformal_upper_bound", "perfect_foresight"])
    ].copy()
    return data[
        [
            (str(hp), int(size)) in valid_pairs
            for hp, size in zip(data.portfolio_type, data.portfolio_size)
        ]
    ].copy()


def portfolio_aggregate(
    portfolio: pd.DataFrame,
    valid_pairs: set[tuple[str, int]],
    group_columns: list[str],
) -> pd.DataFrame:
    required = [
        "portfolio_type",
        "portfolio_size",
        "model",
        "strategy",
        "flexible_fraction",
        "shift_hours",
        "forecast_origins",
        "mean_peak_reduction_fraction",
        "mean_load_factor_change",
        "false_interventions",
        "missed_events",
        "interventions",
    ]
    require_columns(portfolio, required, Path("peak-management summary"))
    data = filter_portfolio_scenario(portfolio, valid_pairs)
    rows = []
    for keys, group in data.groupby(group_columns, observed=True):
        if not isinstance(keys, tuple):
            keys = (keys,)
        weights = pd.to_numeric(group.forecast_origins, errors="coerce").fillna(0).to_numpy(float)
        total_origins = float(weights.sum())
        if total_origins <= 0:
            continue
        row = dict(zip(group_columns, keys))
        row.update(
            {
                "forecast_origins": int(total_origins),
                "peak_reduction_pct": 100.0
                * float(np.average(pd.to_numeric(group.mean_peak_reduction_fraction), weights=weights)),
                "load_factor_change": float(np.average(pd.to_numeric(group.mean_load_factor_change), weights=weights)),
                "false_interventions_per_100_origins": 100.0
                * float(pd.to_numeric(group.false_interventions).sum())
                / total_origins,
                "missed_events_per_100_origins": 100.0
                * float(pd.to_numeric(group.missed_events).sum())
                / total_origins,
                "interventions_per_origin": float(pd.to_numeric(group.interventions).sum())
                / total_origins,
                "valid_portfolio_strata": int(
                    group[["portfolio_type", "portfolio_size"]].drop_duplicates().shape[0]
                ),
            }
        )
        rows.append(row)
    return pd.DataFrame(rows)


def portfolio_tradeoff_with_ci(
    portfolio_summary: pd.DataFrame,
    portfolio_detail: pd.DataFrame,
    valid_pairs: set[tuple[str, int]],
    n_bootstrap: int = 2_000,
    seed: int = 20260924,
) -> pd.DataFrame:
    """Add portfolio-ID cluster bootstrap CIs to Figure 6 point estimates."""
    group_columns = ["portfolio_type", "model", "strategy"]
    point = portfolio_aggregate(portfolio_summary, valid_pairs, group_columns)
    require_columns(
        portfolio_detail,
        [
            "portfolio_id", "portfolio_type", "portfolio_size", "model", "strategy",
            "flexible_fraction", "shift_hours", "peak_reduction_fraction",
            "false_intervention",
        ],
        Path("peak-management origin detail"),
    )
    data = portfolio_detail.copy()
    data["portfolio_size"] = pd.to_numeric(data.portfolio_size, errors="coerce").astype("Int64")
    data = data[
        np.isclose(pd.to_numeric(data.flexible_fraction), 0.20)
        & pd.to_numeric(data.shift_hours).eq(3)
        & data.strategy.isin(["median", "conformal_upper_bound", "perfect_foresight"])
    ].copy()
    data = data[
        [
            (str(hp), int(size)) in valid_pairs
            for hp, size in zip(data.portfolio_type, data.portfolio_size)
        ]
    ]
    clusters = (
        data.groupby([*group_columns, "portfolio_id"], observed=True, sort=False)
        .agg(
            forecast_origins=("peak_reduction_fraction", "size"),
            mean_peak_reduction_fraction=("peak_reduction_fraction", "mean"),
            false_interventions=("false_intervention", "sum"),
        )
        .reset_index()
    )
    rng = np.random.default_rng(seed)
    intervals: list[dict[str, object]] = []
    for keys, group in clusters.groupby(group_columns, observed=True, sort=True):
        group = group.reset_index(drop=True)
        peak_samples = np.empty(n_bootstrap, dtype=float)
        false_samples = np.empty(n_bootstrap, dtype=float)
        for replicate in range(n_bootstrap):
            sampled = group.iloc[rng.integers(0, len(group), size=len(group))]
            weights = pd.to_numeric(sampled.forecast_origins, errors="coerce").fillna(0).to_numpy(float)
            total_origins = float(weights.sum())
            peak_samples[replicate] = 100.0 * float(
                np.average(
                    pd.to_numeric(sampled.mean_peak_reduction_fraction, errors="coerce"),
                    weights=weights,
                )
            )
            false_samples[replicate] = 100.0 * float(
                pd.to_numeric(sampled.false_interventions, errors="coerce").sum()
            ) / total_origins
        intervals.append(
            {
                **dict(zip(group_columns, keys)),
                "peak_ci_low": float(np.quantile(peak_samples, 0.025)),
                "peak_ci_high": float(np.quantile(peak_samples, 0.975)),
                "false_ci_low": float(np.quantile(false_samples, 0.025)),
                "false_ci_high": float(np.quantile(false_samples, 0.975)),
            }
        )
    return point.merge(
        pd.DataFrame(intervals),
        on=group_columns,
        how="left",
        validate="one_to_one",
    )


def build_table_4(portfolio: pd.DataFrame, valid_pairs: set[tuple[str, int]]) -> pd.DataFrame:
    summary = portfolio_aggregate(portfolio, valid_pairs, ["model", "strategy"])
    summary["Model"] = summary.model.map(lambda x: MODEL_LABELS.get(x, x))
    summary["Decision strategy"] = summary.strategy.map(
        {
            "median": "Median forecast",
            "conformal_upper_bound": "Conformal upper bound",
            "perfect_foresight": "Perfect foresight (diagnostic)",
        }
    )
    output = summary[
        [
            "Model",
            "Decision strategy",
            "valid_portfolio_strata",
            "forecast_origins",
            "peak_reduction_pct",
            "false_interventions_per_100_origins",
            "missed_events_per_100_origins",
            "interventions_per_origin",
            "load_factor_change",
        ]
    ].copy()
    output.columns = [
        "Model",
        "Decision strategy",
        "Valid portfolio strata",
        "Forecast origins",
        "Mean peak reduction (%)",
        "False interventions per 100 origins",
        "Missed events per 100 origins",
        "Interventions per origin",
        "Mean load-factor change",
    ]
    return output.sort_values(
        ["Decision strategy", "Mean peak reduction (%)"], ascending=[True, False]
    ).reset_index(drop=True)


def write_tables(tables: list[tuple[str, pd.DataFrame]], table_dir: Path) -> list[Path]:
    table_dir.mkdir(parents=True, exist_ok=True)
    outputs: list[Path] = []
    for stem, table in tables:
        path = table_dir / f"{stem}.csv"
        table.to_csv(path, index=False, encoding="utf-8-sig", float_format="%.3f")
        outputs.append(path)

    workbook = table_dir / "HEAPO_main_tables.xlsx"
    sheet_names = ["Table 1 Cohort", "Table 2 Models", "Table 3 Paired tests", "Table 4 Portfolio"]
    try:
        with pd.ExcelWriter(workbook, engine="openpyxl") as writer:
            for sheet, (_, table) in zip(sheet_names, tables):
                table.to_excel(writer, sheet_name=sheet, index=False)
        from openpyxl import load_workbook
        from openpyxl.styles import Alignment, Font, PatternFill

        book = load_workbook(workbook)
        header_fill = PatternFill("solid", fgColor="1F4E79")
        for sheet in book.worksheets:
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.fill = header_fill
                cell.font = Font(color="FFFFFF", bold=True)
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for column in sheet.columns:
                letter = column[0].column_letter
                width = min(38, max(11, max(len(str(cell.value or "")) for cell in column) + 2))
                sheet.column_dimensions[letter].width = width
                for cell in column[1:]:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
                    if isinstance(cell.value, float):
                        cell.number_format = "0.000"
            sheet.sheet_view.showGridLines = False
        book.save(workbook)
    except ImportError as exc:
        raise ImportError(
            "Excel output requires openpyxl. Install it with: py -3.11 -m pip install openpyxl"
        ) from exc
    outputs.append(workbook)
    return outputs


def draw_box(ax: plt.Axes, xy: tuple[float, float], width: float, height: float, text: str, fill: str) -> None:
    x, y = xy
    patch = FancyBboxPatch(
        (x, y),
        width,
        height,
        boxstyle="round,pad=0.015,rounding_size=0.018",
        transform=ax.transAxes,
        linewidth=1.2,
        edgecolor="#46515C",
        facecolor=fill,
    )
    ax.add_patch(patch)
    ax.text(
        x + width / 2,
        y + height / 2,
        text,
        transform=ax.transAxes,
        ha="center",
        va="center",
        fontsize=9.2,
        linespacing=1.25,
    )


def draw_arrow(ax: plt.Axes, start: tuple[float, float], end: tuple[float, float]) -> None:
    arrow = FancyArrowPatch(
        start,
        end,
        transform=ax.transAxes,
        arrowstyle="-|>",
        mutation_scale=13,
        linewidth=1.1,
        color="#64717D",
        connectionstyle="arc3,rad=0.0",
    )
    ax.add_patch(arrow)


def figure_1(table_1: pd.DataFrame) -> plt.Figure:
    lookup = table_1.set_index("Cohort")
    all_row = lookup.loc["Complete HEAPO dataset"]
    eligible = lookup.loc["Eligible households"]
    train = lookup.loc["Training households"]
    val = lookup.loc["Unseen validation households"]
    test = lookup.loc["Unseen test households"]

    fig, ax = plt.subplots(figsize=(7.4, 4.35))
    ax.axis("off")
    draw_box(
        ax,
        (0.30, 0.77),
        0.40,
        0.15,
        f"HEAPO source population\nn = {int(all_row.Total):,} households\n"
        f"ASHP {int(all_row.ASHP):,} | GSHP {int(all_row.GSHP):,} | unknown {int(all_row['Unknown HP type']):,}",
        "#F2F4F7",
    )
    draw_box(
        ax,
        (0.34, 0.51),
        0.32,
        0.14,
        f"Eligible analysis cohort\nn = {int(eligible.Total)}\n"
        f"ASHP {int(eligible.ASHP)} | GSHP {int(eligible.GSHP)}",
        "#DCEAF7",
    )
    split_specs = [
        (0.03, train, "Model development", "#E5F3EA"),
        (0.365, val, "Model-selection validation\n(LightGBM early stopping)", "#FFF1D6"),
        (0.70, test, "Final unseen test", "#FBE5E2"),
    ]
    for x, row, title, fill in split_specs:
        draw_box(
            ax,
            (x, 0.18),
            0.27,
            0.17,
            f"{title}\nn = {int(row.Total)}\nASHP {int(row.ASHP)} | GSHP {int(row.GSHP)}",
            fill,
        )
    draw_arrow(ax, (0.50, 0.77), (0.50, 0.66))
    draw_arrow(ax, (0.50, 0.51), (0.165, 0.36))
    draw_arrow(ax, (0.50, 0.51), (0.50, 0.36))
    draw_arrow(ax, (0.50, 0.51), (0.835, 0.36))
    return fig


def figure_2(intervals: pd.DataFrame) -> plt.Figure:
    required = ["model", "split", "HP_Type", "horizon", "mae", "mae_ci_low", "mae_ci_high"]
    require_columns(intervals, required, Path("common-index intervals"))
    data = intervals[
        intervals.horizon.isin(HORIZONS)
        & intervals["split"].isin(["seen_test", "unseen_test"])
    ].copy()
    fig, axes = plt.subplots(2, 2, figsize=(7.4, 6.15), sharex=True, sharey=True)
    x = np.arange(len(HORIZONS))
    for index, ((split, hp), ax) in enumerate(zip(STRATUM_ORDER, axes.flat)):
        panel = data[(data["split"] == split) & (data.HP_Type == hp)]
        for model in MODEL_ORDER:
            group = panel[panel.model == model].set_index("horizon").reindex(HORIZONS)
            if group.mae.isna().all():
                continue
            y = group.mae.to_numpy(float)
            lo = group.mae_ci_low.to_numpy(float)
            hi = group.mae_ci_high.to_numpy(float)
            ax.errorbar(
                x,
                y,
                yerr=np.vstack([y - lo, hi - y]),
                color=COLORS[model],
                marker="o" if model != "tft" else "s",
                markersize=3.8,
                linewidth=1.45 if model != "tft" else 1.8,
                linestyle="--" if model == "tft" else "-",
                elinewidth=0.65,
                capsize=1.8,
                capthick=0.65,
                label=MODEL_LABELS[model],
                zorder=3,
            )
        ax.set_ylim(0.0, 0.7)
        ax.set_xticks(x, [HORIZON_LABELS[h] for h in HORIZONS])
        ax.grid(axis="y", color=GRID, linewidth=0.65)
        ax.set_axisbelow(True)
        panel_label(ax, f"{chr(65 + index)}  {STRATUM_LABELS[(split, hp)]}")
        if index % 2 == 0:
            ax.set_ylabel("MAE (kWh per 15 min)")
        if index >= 2:
            ax.set_xlabel("Forecast horizon")
    handles, labels = axes.flat[0].get_legend_handles_labels()
    fig.legend(handles, labels, loc="lower center", ncol=3, frameon=False, bbox_to_anchor=(0.5, -0.01))
    fig.tight_layout(rect=(0, 0.07, 1, 1.0))
    return fig


def figure_3(intervals: pd.DataFrame) -> plt.Figure:
    require_columns(intervals, ["model", "split", "HP_Type", "horizon", "picp"], Path("common-index intervals"))
    data = intervals[
        intervals.horizon.isin(HORIZONS)
        & intervals["split"].isin(["seen_test", "unseen_test"])
    ].copy()
    deviations = pd.to_numeric(data.picp, errors="coerce") - 0.80
    max_abs = max(0.10, float(np.nanmax(np.abs(deviations))))
    norm = mcolors.TwoSlopeNorm(vmin=-max_abs, vcenter=0.0, vmax=max_abs)
    cmap = plt.get_cmap("RdBu")
    fig, axes = plt.subplots(2, 2, figsize=(7.4, 5.15), sharex=True, sharey=True)
    image = None
    for index, ((split, hp), ax) in enumerate(zip(STRATUM_ORDER, axes.flat)):
        panel = data[(data["split"] == split) & (data.HP_Type == hp)]
        matrix = (
            panel.pivot(index="model", columns="horizon", values="picp")
            .reindex(index=MODEL_ORDER, columns=HORIZONS)
            .astype(float)
        )
        values = matrix.to_numpy() - 0.80
        image = ax.pcolormesh(
            np.arange(values.shape[1] + 1) - 0.5,
            np.arange(values.shape[0] + 1) - 0.5,
            values,
            cmap=cmap,
            norm=norm,
            shading="flat",
            edgecolors="white",
            linewidth=0.7,
            rasterized=False,
        )
        for row in range(matrix.shape[0]):
            for col in range(matrix.shape[1]):
                value = matrix.iat[row, col]
                if not np.isfinite(value):
                    continue
                background = norm(value - 0.80)
                color = "white" if background < 0.20 or background > 0.80 else "#111111"
                ax.text(col, row, f"{value:.2f}", ha="center", va="center", fontsize=7.7, color=color)
        households = panel["n_households"].max() if "n_households" in panel.columns else np.nan
        count_text = f"; n = {int(households)}" if np.isfinite(households) else ""
        ax.set_xticks(range(len(HORIZONS)), [HORIZON_LABELS[h] for h in HORIZONS])
        ax.set_yticks(range(len(MODEL_ORDER)), [MODEL_LABELS[m] for m in MODEL_ORDER])
        ax.set_ylim(-0.5, len(MODEL_ORDER) - 0.5)
        panel_header(
            ax,
            f"{chr(65 + index)}. {STRATUM_LABELS[(split, hp)]}{count_text}",
        )
        for spine in ax.spines.values():
            spine.set_visible(False)
    assert image is not None
    fig.subplots_adjust(
        left=0.18,
        right=0.98,
        top=0.93,
        bottom=0.155,
        hspace=0.34,
        wspace=0.16,
    )
    fig.text(0.58, 0.105, "Forecast horizon", ha="center", va="center", fontsize=9.5)
    colorbar_axis = fig.add_axes([0.24, 0.028, 0.55, 0.026])
    colorbar = fig.colorbar(image, cax=colorbar_axis, orientation="horizontal")
    if colorbar.solids is not None:
        colorbar.solids.set_rasterized(False)
    colorbar.set_label("PICP deviation from the nominal 0.80 coverage")
    return fig


def forest_panel(
    ax: plt.Axes,
    frame: pd.DataFrame,
    candidate_label: str,
    color: str,
    panel: str,
) -> None:
    data = frame[frame.confirmatory_holm_primary.astype(bool)].copy()
    data["horizon_order"] = data.horizon.map({h: i for i, h in enumerate(HORIZONS)})
    data["comparator_order"] = data.comparator.map({m: i for i, m in enumerate(MODEL_ORDER)})
    data = data.sort_values(["horizon_order", "comparator_order"], ascending=[False, True]).reset_index(drop=True)
    y = np.arange(len(data))
    for idx, row in data.iterrows():
        robust = bool(row.confirmatory_holm_nested)
        ax.plot([row.ci_low_primary, row.ci_high_primary], [idx, idx], color=color, linewidth=1.4, zorder=2)
        ax.scatter(
            row.mae_delta_primary,
            idx,
            s=38,
            marker="o",
            facecolor=color if robust else "white",
            edgecolor=color,
            linewidth=1.4,
            zorder=3,
        )
    labels = [
        f"{HORIZON_LABELS[int(row.horizon)]} vs {MODEL_LABELS.get(row.comparator, row.comparator)}"
        for _, row in data.iterrows()
    ]
    ax.set_yticks(y, labels)
    ax.axvline(0, color="#20252A", linewidth=1.0)
    ax.grid(axis="x", color=GRID, linewidth=0.65)
    ax.set_axisbelow(True)
    panel_header(ax, f"{panel}. {candidate_label}")
    if len(data):
        lo = float(data.ci_low_primary.min())
        hi = float(data.ci_high_primary.max())
        span = max(hi - lo, 0.02)
        margin = 0.10 * span
        null_margin = 0.05 * span
        ax.set_xlim(min(lo - margin, -null_margin), max(hi + margin, null_margin))


def figure_4(tabular: pd.DataFrame, tft: pd.DataFrame) -> plt.Figure:
    fig, axes = plt.subplots(
        1,
        2,
        figsize=(7.4, 4.35),
        gridspec_kw={"width_ratios": [0.92, 1.28]},
    )
    forest_panel(axes[0], tabular, "LightGBM comparisons", COLORS["lightgbm"], "A")
    forest_panel(axes[1], tft, "TFT comparisons", COLORS["tft"], "B")
    legend = [
        Line2D([0], [0], marker="o", color="#444444", markerfacecolor="#444444", linewidth=0, label="Confirmed in both designs"),
        Line2D([0], [0], marker="o", color="#444444", markerfacecolor="white", linewidth=0, label="Primary design only"),
    ]
    fig.legend(
        handles=legend,
        loc="lower center",
        ncol=2,
        frameon=False,
        bbox_to_anchor=(0.5, 0.005),
    )
    fig.supxlabel(
        "Paired ΔMAE (candidate − comparator)", x=0.5, y=0.105, fontsize=9.5
    )
    fig.text(
        0.5,
        0.061,
        "Candidate better  ←                    0                    →  Candidate worse",
        ha="center",
        fontsize=8.5,
        color="#4F5B66",
    )
    fig.tight_layout(rect=(0, 0.145, 1, 0.95), w_pad=2.0)
    return fig


def pretty_feature(name: str) -> str:
    exact = {
        "load_on_fraction_24h": "On fraction (24 h)",
        "load_roll_mean_96": "Rolling load mean (24 h)",
        "load_roll_mean_48": "Rolling load mean (12 h)",
        "load_roll_mean_24": "Rolling load mean (6 h)",
        "load_roll_mean_12": "Rolling load mean (3 h)",
        "load_roll_mean_4": "Rolling load mean (1 h)",
        "load_day_base": "Daily base load",
        "load_day_max": "Daily maximum load",
        "load_day_min": "Daily minimum load",
        "load_lag_1": "Load lag (15 min)",
        "load_lag_4": "Load lag (1 h)",
        "load_lag_24": "Load lag (6 h)",
        "future_weather__Temperature_avg_hourly": "Outdoor temperature",
        "future_weather__Sunshine_duration_hourly": "Sunshine duration",
        "future__cal_hour_sin": "Hour-of-day sine",
        "future__cal_hour_cos": "Hour-of-day cosine",
    }
    if name in exact:
        return exact[name]
    cleaned = name
    for prefix in ("future_weather__", "future__cal_", "static__Survey_", "static__", "load_"):
        cleaned = cleaned.replace(prefix, "")
    return cleaned.replace("_", " ").strip().title()


def normalize_shap(shap: pd.DataFrame) -> pd.DataFrame:
    frame = shap.copy()
    rename = {}
    if "Feature" in frame.columns:
        rename["Feature"] = "feature"
    if "mean_abs_SHAP" in frame.columns:
        rename["mean_abs_SHAP"] = "mean_abs_shap"
    if "mean_absolute_shap" in frame.columns:
        rename["mean_absolute_shap"] = "mean_abs_shap"
    if "hp_type" in frame.columns:
        rename["hp_type"] = "HP_Type"
    frame = frame.rename(columns=rename)
    if "feature_label" in frame.columns:
        frame["feature"] = frame["feature_label"].astype(str)
    if "mean_abs_shap" not in frame.columns and {"feature", "shap_value"}.issubset(frame.columns):
        group_cols = ["feature"] + (["HP_Type"] if "HP_Type" in frame.columns else [])
        frame = (
            frame.assign(_abs=pd.to_numeric(frame.shap_value, errors="coerce").abs())
            .groupby(group_cols, observed=True)._abs.mean()
            .rename("mean_abs_shap")
            .reset_index()
        )
    require_columns(frame, ["feature", "mean_abs_shap"], Path("SHAP importance"))
    if "HP_Type" not in frame.columns:
        frame["HP_Type"] = "All"
    frame["feature"] = frame.feature.astype(str)
    frame["HP_Type"] = frame.HP_Type.astype(str)
    frame["mean_abs_shap"] = pd.to_numeric(frame.mean_abs_shap, errors="coerce")
    return frame.dropna(subset=["mean_abs_shap"])


def figure_5(shap: pd.DataFrame, top_n: int = 12) -> plt.Figure:
    data = normalize_shap(shap)
    ranking = data.groupby("feature", observed=True).mean_abs_shap.mean().nlargest(top_n)
    features = ranking.sort_values().index.tolist()
    pivot = (
        data[data.feature.isin(features)]
        .pivot_table(index="feature", columns="HP_Type", values="mean_abs_shap", aggfunc="mean")
        .reindex(features)
    )
    fig, ax = plt.subplots(figsize=(7.4, 5.1))
    y = np.arange(len(features))
    if {"ASHP", "GSHP"}.issubset(pivot.columns):
        ashp = pivot.ASHP.to_numpy(float)
        gshp = pivot.GSHP.to_numpy(float)
        for idx in range(len(features)):
            ax.plot([ashp[idx], gshp[idx]], [idx, idx], color="#B9C1C9", linewidth=1.4, zorder=1)
        ax.scatter(ashp, y, color="#0072B2", s=34, label="ASHP", zorder=3)
        ax.scatter(gshp, y, color="#D55E00", s=34, marker="s", label="GSHP", zorder=3)
    else:
        column = pivot.columns[0]
        ax.barh(y, pivot[column].to_numpy(float), color="#0072B2", height=0.64)
    ax.set_yticks(y, [pretty_feature(feature) for feature in features])
    ax.set_xlabel("Mean absolute SHAP value (kWh per 15 min)")
    ax.grid(axis="x", color=GRID, linewidth=0.65)
    ax.set_axisbelow(True)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    if {"ASHP", "GSHP"}.issubset(pivot.columns):
        ax.legend(frameon=False, loc="lower right")
    fig.tight_layout()
    return fig


def figure_6(
    portfolio: pd.DataFrame,
    portfolio_detail: pd.DataFrame,
    valid_pairs: set[tuple[str, int]],
) -> plt.Figure:
    summary = portfolio_tradeoff_with_ci(portfolio, portfolio_detail, valid_pairs)
    types = [hp for hp in ["ASHP", "GSHP", "mixed"] if hp in set(summary.portfolio_type)]
    fig, axes = plt.subplots(1, len(types), figsize=(7.4, 3.45), sharex=True, sharey=True)
    if len(types) == 1:
        axes = np.array([axes])
    markers = {"median": "o", "conformal_upper_bound": "^"}
    for index, (hp, ax) in enumerate(zip(types, axes.flat)):
        panel = summary[summary.portfolio_type == hp]
        for model in [m for m in MODEL_ORDER if m != "tft"]:
            points = panel[panel.model == model]
            if points.empty:
                continue
            points = points.set_index("strategy")
            if {"median", "conformal_upper_bound"}.issubset(points.index):
                ax.plot(
                    points.loc[["median", "conformal_upper_bound"], "false_interventions_per_100_origins"],
                    points.loc[["median", "conformal_upper_bound"], "peak_reduction_pct"],
                    color=COLORS[model],
                    linewidth=0.9,
                    alpha=0.55,
                    zorder=1,
                )
            for strategy, row in points.iterrows():
                x_value = float(row.false_interventions_per_100_origins)
                y_value = float(row.peak_reduction_pct)
                ax.errorbar(
                    x_value,
                    y_value,
                    xerr=np.array(
                        [[max(0.0, x_value - float(row.false_ci_low))],
                         [max(0.0, float(row.false_ci_high) - x_value)]]
                    ),
                    yerr=np.array(
                        [[max(0.0, y_value - float(row.peak_ci_low))],
                         [max(0.0, float(row.peak_ci_high) - y_value)]]
                    ),
                    color=COLORS[model],
                    marker=markers[strategy],
                    markersize=5.8,
                    markeredgecolor="white",
                    markeredgewidth=0.5,
                    elinewidth=0.65,
                    capsize=1.8,
                    linestyle="none",
                    alpha=0.90,
                    zorder=3,
                )
        oracle = panel[
            (panel.model == "perfect_foresight")
            & (panel.strategy == "perfect_foresight")
        ]
        for _, row in oracle.iterrows():
            x_value = float(row.false_interventions_per_100_origins)
            y_value = float(row.peak_reduction_pct)
            ax.errorbar(
                x_value,
                y_value,
                xerr=np.array(
                    [[max(0.0, x_value - float(row.false_ci_low))],
                     [max(0.0, float(row.false_ci_high) - x_value)]]
                ),
                yerr=np.array(
                    [[max(0.0, y_value - float(row.peak_ci_low))],
                     [max(0.0, float(row.peak_ci_high) - y_value)]]
                ),
                color=COLORS["perfect_foresight"],
                marker="*",
                markersize=8.0,
                markeredgecolor="white",
                markeredgewidth=0.5,
                elinewidth=0.75,
                capsize=1.8,
                linestyle="none",
                zorder=4,
            )
        label = "Mixed" if hp == "mixed" else hp
        ax.set_xlabel("False interventions\nper 100 forecast origins")
        ax.axhline(0.0, color="#4F5B66", linewidth=0.8, zorder=0)
        ax.grid(color=GRID, linewidth=0.65)
        ax.set_axisbelow(True)
        panel_header(ax, f"{chr(65 + index)}. {label} portfolios")
        if index == 0:
            ax.set_ylabel("Mean peak reduction (%)")
    model_handles = [
        Line2D([0], [0], marker="o", color=COLORS[m], linewidth=1.2, label=MODEL_LABELS[m], markersize=5)
        for m in MODEL_ORDER
        if m != "tft"
    ]
    model_handles.append(
        Line2D(
            [0], [0], marker="*", color=COLORS["perfect_foresight"],
            linewidth=0, label="Perfect foresight", markersize=7
        )
    )
    strategy_handles = [
        Line2D([0], [0], marker="o", color="#333333", linewidth=0, label="Median forecast", markersize=5),
        Line2D([0], [0], marker="^", color="#333333", linewidth=0, label="Conformal upper bound", markersize=6),
    ]
    fig.legend(
        handles=model_handles + strategy_handles,
        loc="lower center",
        ncol=4,
        frameon=False,
        bbox_to_anchor=(0.5, 0.005),
    )
    fig.tight_layout(rect=(0, 0.165, 1, 0.95), w_pad=1.1)
    return fig


def write_captions(path: Path, intervals: pd.DataFrame) -> None:
    counts = []
    for split, hp in STRATUM_ORDER:
        panel = intervals[(intervals["split"] == split) & (intervals.HP_Type == hp)]
        if panel.empty or not {"n_households", "n_rows"}.issubset(panel.columns):
            continue
        counts.append(
            f"{STRATUM_LABELS[(split, hp)]}: n = {int(panel.n_households.max())} households "
            f"and {int(panel.n_rows.max())} common-index forecast origins per horizon"
        )
    count_text = "; ".join(counts)
    text = f"""# HEAPO main-paper figure and table captions

## Tables

**Table 1.** Study population and household-level data partitions. The seen-test row reports independent households contributing to the final within-household evaluation; the confirmatory threshold was 20 households.

**Table 2.** Common-index model summary across the four operational forecast horizons. Metrics were weighted by the number of common forecast-origin observations. Lower MAE, approximate CRPS, interval width and weighted interval score (WIS) indicate better performance conditional on calibration; nominal prediction-interval coverage was 0.80.

**Table 3.** Confirmatory paired MAE comparisons from the primary household-cluster bootstrap and the nested household/local-day sensitivity analysis. Delta MAE was calculated as candidate minus comparator; negative values favour the candidate. Holm correction was applied within each pre-specified candidate/comparator family.

**Table 4.** Leakage-resistant portfolio peak-management outcomes under the pre-specified 20% flexibility and 3 h forward-shifting scenario. Forecast-driven source, destination and acceptance decisions were frozen from origin-available information; realised demand was used only for ex post scoring. The non-deployable perfect-foresight row applies the identical dispatch constraints to realised future load and separates forecast loss from dispatch limitations. Values were weighted by forecast origins across calibration-valid portfolio strata.

## Figures

**Figure 1.** HEAPO study population and leakage-resistant evaluation design. Model parameters were fitted using development-household targets ending by the training cutoff; household-disjoint validation origins and their complete 96-step targets also ended by that cutoff and were used only for LightGBM early-stopping/model-selection monitoring. The final unseen-test group was opened only for locked generalization evaluation. Primary inference resampled complete households; the sensitivity analysis additionally resampled local calendar days within each sampled household occurrence. Confirmatory inference required at least 20 independent households.

**Figure 2.** Multi-horizon MAE on the common TFT/tabular evaluation index. Lines show model-specific MAE at 15 min, 1 h, 6 h and 24 h; thin error bars show 95% household-cluster bootstrap confidence intervals. Sample sizes (independent households/common-index forecast origins per horizon) were: {count_text}.

**Figure 3.** Prediction-interval coverage on the common evaluation index. Cell values are PICP estimates and colours encode deviation from the nominal 0.80 coverage; blue and red represent over- and under-coverage, respectively. Household counts are shown in the panel identifiers. The heatmap cells and text are rendered as vector elements in the PDF.

**Figure 4.** Confirmatory paired MAE effects in the seen-ASHP stratum. Points show candidate-minus-comparator MAE differences and horizontal lines show 95% household-cluster bootstrap confidence intervals. Filled points remained confirmatory in the nested household/local-day sensitivity analysis; open points were significant only in the primary design.

**Figure 5.** Global LightGBM feature importance for ASHP and GSHP observations. Only the 12 highest-ranked features are shown. Importance was quantified as the mean absolute SHAP value and averaged over a seeded sample of held-out seen- and unseen-test origin–horizon observations within each heat-pump type (maximum 5,000 observations per type) for the LightGBM mean model; features were then ranked by their mean importance across heat-pump types.

**Figure 6.** Peak-reduction and false-intervention trade-off for calibration-valid virtual portfolios under 20% flexible load and a 3 h forward-shifting window. Median forecasts and conformal upper bounds were used as deployable decision signals; the star denotes the non-deployable perfect-foresight diagnostic under identical dispatch constraints. Values were aggregated across seen and unseen cohorts and valid portfolio sizes using forecast origins as weights. Horizontal and vertical error bars show 95% portfolio-ID cluster-bootstrap confidence intervals. The portfolio-calibration sensitivity analysis is reported in Supplementary Table S1.
"""
    path.write_text(text, encoding="utf-8")


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def validate_outputs(outputs: Sequence[Path], expected_figures: int = 6, expected_tables: int = 4) -> None:
    missing = [path for path in outputs if not path.exists() or path.stat().st_size == 0]
    if missing:
        raise RuntimeError(f"Output validation failed; empty or missing files: {missing}")
    pngs = [path for path in outputs if path.suffix.lower() == ".png"]
    rasters = [path for path in outputs if path.suffix.lower() in {".png", ".tif", ".tiff"}]
    csvs = [path for path in outputs if path.suffix.lower() == ".csv"]
    if len(pngs) != expected_figures:
        raise RuntimeError(f"Expected {expected_figures} PNG figures, found {len(pngs)}")
    if len(csvs) != expected_tables:
        raise RuntimeError(f"Expected {expected_tables} CSV tables, found {len(csvs)}")
    from PIL import Image

    for path in rasters:
        with Image.open(path) as image:
            width, height = image.size
            dpi = image.info.get("dpi", (0.0, 0.0))
            image.verify()
        if min(width, height) < 1000:
            raise RuntimeError(
                f"Figure resolution is unexpectedly low: {path} -> {(width, height)}"
            )
        if min(float(dpi[0]), float(dpi[1])) < 590:
            raise RuntimeError(f"Figure DPI is below 600-dpi tolerance: {path} -> {dpi}")


def main() -> None:
    args = parse_args()
    configure_style()
    run_dir = args.run_dir.resolve()
    if not run_dir.exists():
        raise FileNotFoundError(f"Run directory does not exist: {run_dir}")
    validation_path = run_dir / "validation_report.json"
    if not validation_path.exists():
        raise RuntimeError("validation_report.json is missing; complete the full run first")
    validation = json.loads(validation_path.read_text(encoding="utf-8"))
    if not bool(validation.get("passed", False)):
        failed = [
            item.get("check", "unknown")
            for item in validation.get("checks", [])
            if not bool(item.get("passed", False))
        ]
        raise RuntimeError(
            "Manuscript export is blocked because validation did not pass: "
            + ", ".join(failed)
        )
    out_dir = (args.out_dir or run_dir / "manuscript_outputs").resolve()
    table_dir = out_dir / "tables"
    figure_dir = out_dir / "figures"
    out_dir.mkdir(parents=True, exist_ok=True)

    sources = resolve_sources(run_dir, args.statistics_dir, args.audit_dir, args.shap_file)
    print("Resolved source files:")
    for name, path in sources.__dict__.items():
        print(f"  {name}: {path}")

    audit = read_frame(sources.audit)
    split_manifest = json.loads(sources.split_manifest.read_text(encoding="utf-8"))
    common_metrics = read_frame(sources.common_metrics)
    intervals = read_frame(sources.common_intervals)
    tabular_primary = read_frame(sources.tabular_primary)
    tabular_nested = read_frame(sources.tabular_nested)
    tft_primary = read_frame(sources.tft_primary)
    tft_nested = read_frame(sources.tft_nested)
    shap = read_frame(sources.shap)
    portfolio = read_frame(sources.portfolio_summary)
    portfolio_detail = read_frame(sources.portfolio_detail)
    corrections = read_frame(sources.portfolio_corrections)

    tabular = merge_paired(tabular_primary, tabular_nested)
    tft = merge_paired(tft_primary, tft_nested)
    table_1 = build_table_1(audit, split_manifest, intervals)
    table_2, _weighted = weighted_model_summary(common_metrics)
    table_3 = build_table_3(tabular, tft)
    valid_pairs = valid_portfolio_pairs(corrections)
    table_4 = build_table_4(portfolio, valid_pairs)

    table_specs = [
        ("Table_1_study_population", table_1),
        ("Table_2_common_index_model_summary", table_2),
        ("Table_3_confirmatory_paired_comparisons", table_3),
        ("Table_4_portfolio_peak_management", table_4),
    ]
    outputs = write_tables(table_specs, table_dir)

    figures = [
        ("Figure_1_study_design", figure_1(table_1)),
        ("Figure_2_multihorizon_MAE", figure_2(intervals)),
        ("Figure_3_interval_calibration", figure_3(intervals)),
        ("Figure_4_confirmatory_forest_plot", figure_4(tabular, tft)),
        ("Figure_5_SHAP_importance", figure_5(shap)),
        (
            "Figure_6_portfolio_tradeoff",
            figure_6(portfolio, portfolio_detail, valid_pairs),
        ),
    ]
    for stem, fig in figures:
        outputs.extend(save_figure(fig, figure_dir / stem, args.dpi, args.tiff))

    captions = out_dir / "MANUSCRIPT_CAPTIONS.md"
    write_captions(captions, intervals)
    outputs.append(captions)

    validate_outputs(outputs)
    manifest = {
        "run_dir": str(run_dir),
        "output_dir": str(out_dir),
        "figure_dpi": args.dpi,
        "tiff_enabled": bool(args.tiff),
        "sources": {name: str(path) for name, path in sources.__dict__.items()},
        "valid_portfolio_pairs": sorted([list(item) for item in valid_pairs]),
        "outputs": [
            {
                "path": str(path),
                "bytes": path.stat().st_size,
                "sha256": sha256(path),
            }
            for path in outputs
        ],
    }
    manifest_path = out_dir / "manuscript_output_manifest.json"
    manifest_path.write_text(json.dumps(manifest, indent=2, ensure_ascii=False), encoding="utf-8")

    print("\nCompleted successfully.")
    print(f"Tables:  {table_dir}")
    print(f"Figures: {figure_dir}")
    print(f"Captions: {captions}")
    print(f"Manifest: {manifest_path}")


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        raise
